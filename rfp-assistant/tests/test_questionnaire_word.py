import io

import docx

from rfp.export import write_back_word
from rfp.questionnaire import detect_word_questions, is_word_question

from .helpers import make_docx

DOC = make_docx(
    [
        ("Title", "Security Questionnaire"),
        ("Normal", "Please answer all questions below."),
        ("Heading 1", "General"),
        ("Normal", "1. Where is customer data hosted?"),
        ("Normal", "Describe your backup policy."),
        ("Normal", "Thank you for your time."),
    ],
    table=[["Question", "Answer"], ["Do you support SSO?", ""], ["Is MFA available?", ""]],
)


def test_detects_paragraph_and_table_questions():
    questions = detect_word_questions(DOC)
    assert [q.text for q in questions] == [
        "1. Where is customer data hosted?",
        "Describe your backup policy.",
        "Do you support SSO?",
        "Is MFA available?",
    ]
    assert [q.kind for q in questions] == ["paragraph", "paragraph", "table", "table"]
    assert questions[2].location == {"table": 0, "row": 1, "col": 0, "answer_col": 1}


def test_question_heuristic():
    assert is_word_question("What is your SLA?")
    assert is_word_question("Please describe your onboarding process.")
    assert is_word_question("3.1 Provide details of your disaster recovery plan")
    assert not is_word_question("Thank you for your time.")
    assert not is_word_question("Security")


def test_write_back_word():
    questions = detect_word_questions(DOC)
    answers = ["In the EU.", "No answer found", "Yes, SAML 2.0.", "Yes."]
    out = write_back_word(DOC, list(zip(questions, answers)), skip_no_answer=True)
    d = docx.Document(io.BytesIO(out))
    texts = [p.text for p in d.paragraphs]
    i = texts.index("1. Where is customer data hosted?")
    assert texts[i + 1] == "Answer: In the EU."
    assert "Answer: No answer found" not in texts
    assert d.tables[0].cell(1, 1).text == "Yes, SAML 2.0."
    assert d.tables[0].cell(2, 1).text == "Yes."
