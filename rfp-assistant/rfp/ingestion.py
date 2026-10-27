"""Turn uploaded PDF, Word and Excel files into chunks with metadata."""
from __future__ import annotations

import hashlib
import io
import re
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".xlsm", ".txt", ".md"}


@dataclass
class Section:
    """A piece of a document with a human-readable location (page, sheet, ...)."""

    location: str
    text: str


@dataclass
class Chunk:
    id: str
    text: str
    metadata: dict = field(default_factory=dict)


def document_id(filename: str) -> str:
    """Stable id per file name, so re-uploading a file replaces the old version."""
    return hashlib.sha1(filename.strip().lower().encode()).hexdigest()[:16]


# ---------------------------------------------------------------- extraction


def extract_sections(filename: str, data: bytes) -> list[Section]:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return _extract_pdf(data)
    if ext == ".docx":
        return _extract_docx(data)
    if ext in (".xlsx", ".xlsm"):
        return _extract_xlsx(data)
    if ext in (".txt", ".md"):
        return [Section("document", data.decode("utf-8", errors="replace"))]
    raise ValueError(f"Unsupported file type: {ext or filename}")


def _extract_pdf(data: bytes) -> list[Section]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    sections = []
    for number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            sections.append(Section(f"page {number}", text))
    return sections


def _extract_docx(data: bytes) -> list[Section]:
    """Group paragraphs under their nearest heading; tables become 'A | B' rows."""
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(io.BytesIO(data))
    sections: list[Section] = []
    heading = "document"
    lines: list[str] = []
    has_body = False  # headings with no text under them merge into the next section

    def flush() -> None:
        text = "\n".join(lines).strip()
        if text:
            sections.append(Section(f"section '{heading}'" if heading != "document" else heading, text))
        lines.clear()

    for block in document.element.body.iterchildren():
        tag = block.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(block, document)
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name if para.style is not None else "") or ""
            if style.lower().startswith(("heading", "title")):
                if has_body:
                    flush()
                    has_body = False
                heading = text[:80]
            else:
                has_body = True
            lines.append(text)
        elif tag == "tbl":
            table = Table(block, document)
            for row in table.rows:
                cells = _unique_cells(cell.text.strip() for cell in row.cells)
                if any(cells):
                    lines.append(" | ".join(cells))
                    has_body = True
    flush()
    return sections


def _unique_cells(cells) -> list[str]:
    """Merged Word cells repeat their text; keep each value once per row."""
    seen: list[str] = []
    for cell in cells:
        if not seen or seen[-1] != cell:
            seen.append(cell)
    return seen


def _extract_xlsx(data: bytes, rows_per_section: int = 25) -> list[Section]:
    """Each row becomes 'Header: value; Header: value' so meaning survives chunking."""
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sections: list[Section] = []
    for sheet in workbook.worksheets:
        rows = [
            (index, [_cell_text(v) for v in row])
            for index, row in enumerate(sheet.iter_rows(values_only=True), start=1)
        ]
        rows = [(i, r) for i, r in rows if any(r)]
        if not rows:
            continue
        first = rows[0][1]
        if sum(1 for c in first if c) > 1 and len(rows) > 1:
            header, body = first, rows[1:]
        else:
            header, body = [], rows
        for start in range(0, len(body), rows_per_section):
            group = body[start : start + rows_per_section]
            lines = [_row_text(header, values) for _, values in group]
            first, last = group[0][0], group[-1][0]
            sections.append(Section(f"sheet '{sheet.title}' rows {first}-{last}", "\n".join(lines)))
    workbook.close()
    return sections


def _cell_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _row_text(header: list[str], values: list[str]) -> str:
    parts = []
    for i, value in enumerate(values):
        if not value:
            continue
        label = header[i] if i < len(header) and header[i] else ""
        parts.append(f"{label}: {value}" if label else value)
    return "; ".join(parts)


# ------------------------------------------------------------------ chunking


def chunk_text(text: str, size: int = 1200, overlap: int = 200) -> list[str]:
    """Split on paragraph and sentence boundaries into chunks of about `size` chars."""
    text = re.sub(r"[ \t]+", " ", text).strip()
    if len(text) <= size:
        return [text] if text else []

    pieces: list[str] = []
    for paragraph in re.split(r"\n\s*\n|\n", text):
        paragraph = paragraph.strip()
        if not paragraph:
            continue
        if len(paragraph) <= size:
            pieces.append(paragraph)
        else:
            pieces.extend(_split_long(paragraph, size))

    chunks: list[str] = []
    current = ""
    for piece in pieces:
        if current and len(current) + len(piece) + 1 > size:
            chunks.append(current)
            tail = current[-overlap:] if overlap else ""
            # start the overlap at a word boundary
            tail = tail[tail.find(" ") + 1 :] if " " in tail else tail
            current = f"{tail} {piece}".strip() if tail else piece
        else:
            current = f"{current}\n{piece}" if current else piece
    if current:
        chunks.append(current)
    return chunks


def _split_long(paragraph: str, size: int) -> list[str]:
    sentences = re.split(r"(?<=[.!?])\s+", paragraph)
    out: list[str] = []
    for sentence in sentences:
        while len(sentence) > size:
            out.append(sentence[:size])
            sentence = sentence[size:]
        if sentence:
            out.append(sentence)
    return out


def ingest_file(
    filename: str,
    data: bytes,
    chunk_size: int = 1200,
    chunk_overlap: int = 200,
    upload_date: str | None = None,
) -> list[Chunk]:
    filename = Path(filename).name
    doc_id = document_id(filename)
    upload_date = upload_date or datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    chunks: list[Chunk] = []
    for section in extract_sections(filename, data):
        for piece in chunk_text(section.text, chunk_size, chunk_overlap):
            index = len(chunks)
            chunks.append(
                Chunk(
                    id=f"{doc_id}-{index}",
                    text=piece,
                    metadata={
                        "doc_id": doc_id,
                        "source": filename,
                        "location": section.location,
                        "upload_date": upload_date,
                        "chunk_index": index,
                    },
                )
            )
    return chunks
