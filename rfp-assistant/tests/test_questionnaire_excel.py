import io

from openpyxl import load_workbook

from rfp.export import review_table_to_excel, write_back_excel
from rfp.questionnaire import ExcelLayout, detect_excel_layout, extract_excel_questions

from .helpers import make_xlsx

STANDARD = make_xlsx({
    "Cover": [["Acme RFP 2026"], ["Please answer every question."]],
    "Security": [
        ["Vendor Requirements"],
        [],
        ["ID", "Category", "Question", "Vendor Response", "Comments"],
        [1, "Security", "Do you encrypt data at rest?", None, None],
        [2, "Security", "Describe your incident response process.", None, None],
        [3, "Ops", "What is your uptime SLA?", None, None],
    ],
})


def test_detects_header_question_and_answer_columns():
    layout = detect_excel_layout(STANDARD)
    assert layout == ExcelLayout(sheet="Security", header_row=3, question_col=3, answer_col=4)


def test_extracts_questions_with_rows():
    layout = detect_excel_layout(STANDARD)
    questions = extract_excel_questions(STANDARD, layout)
    assert [q.text for q in questions][0] == "Do you encrypt data at rest?"
    assert [q.location["row"] for q in questions] == [4, 5, 6]
    assert questions[0].id == "Security!C4"


def test_headerless_sheet_uses_content_and_first_empty_column():
    data = make_xlsx({"Sheet1": [
        ["1", "Do you support SSO?"],
        ["2", "Is there an API for exporting data?"],
        ["3", "Where is customer data hosted?"],
    ]})
    layout = detect_excel_layout(data)
    assert (layout.header_row, layout.question_col, layout.answer_col) == (0, 2, 3)
    assert len(extract_excel_questions(data, layout)) == 3


def test_write_back_keeps_other_sheets_and_fills_answers():
    layout = detect_excel_layout(STANDARD)
    out = write_back_excel(STANDARD, layout, {4: "Yes, AES-256.", 5: "No answer found"},
                           skip_no_answer=True)
    wb = load_workbook(io.BytesIO(out))
    ws = wb["Security"]
    assert ws["D4"].value == "Yes, AES-256."
    assert ws["D5"].value is None
    assert ws["C4"].value == "Do you encrypt data at rest?"
    assert wb["Cover"]["A1"].value == "Acme RFP 2026"


def test_write_back_adds_header_when_answer_column_is_new():
    data = make_xlsx({"Q": [["#", "Question"], [1, "Do you have SOC 2?"]]})
    layout = detect_excel_layout(data)
    out = write_back_excel(data, layout, {2: "Yes."})
    ws = load_workbook(io.BytesIO(out))["Q"]
    assert ws["C1"].value == "Answer" and ws["C2"].value == "Yes."


def test_review_table_export():
    import pandas as pd

    table = pd.DataFrame([{"Question": "Q?", "Draft answer": "A.", "Sources": "a.pdf (page 1)",
                           "Confidence": "high", "Status": "Approved", "_qid": "x"}])
    ws = load_workbook(io.BytesIO(review_table_to_excel(table)))["Answers"]
    assert [c.value for c in ws[1]] == ["Question", "Draft answer", "Sources", "Confidence", "Status"]
    assert ws["E2"].value == "Approved"
