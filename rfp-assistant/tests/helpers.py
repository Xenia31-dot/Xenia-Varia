"""Builders for small test files and a fake Claude client."""
from __future__ import annotations

import io
import json
from types import SimpleNamespace


def make_pdf(pages: list[str]) -> bytes:
    """A minimal valid PDF with one line of text per page (no extra deps)."""
    objects: list[bytes] = []
    page_ids = []
    font_id = 3 + 2 * len(pages)
    for i, text in enumerate(pages):
        page_id, content_id = 3 + 2 * i, 4 + 2 * i
        page_ids.append(page_id)
        safe = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = f"BT /F1 12 Tf 50 750 Td ({safe}) Tj ET".encode()
        objects.append(
            f"{page_id} 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Contents {content_id} 0 R /Resources << /Font << /F1 {font_id} 0 R >> >> >> endobj\n".encode()
        )
        objects.append(
            f"{content_id} 0 obj << /Length {len(stream)} >> stream\n".encode()
            + stream + b"\nendstream endobj\n"
        )
    kids = " ".join(f"{p} 0 R" for p in page_ids)
    head = [
        b"1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj\n",
        f"2 0 obj << /Type /Pages /Kids [{kids}] /Count {len(pages)} >> endobj\n".encode(),
    ]
    tail = [f"{font_id} 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj\n".encode()]
    body = b"%PDF-1.4\n"
    offsets = []
    for obj in head + objects + tail:
        offsets.append(len(body))
        body += obj
    xref = len(body)
    body += f"xref\n0 {len(offsets) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        body += f"{off:010d} 00000 n \n".encode()
    body += f"trailer << /Size {len(offsets) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return body


def make_docx(paragraphs: list[tuple[str, str]], table: list[list[str]] | None = None) -> bytes:
    """paragraphs: (style, text) pairs, style e.g. 'Heading 1' or 'Normal'."""
    import docx

    document = docx.Document()
    for style, text in paragraphs:
        document.add_paragraph(text, style=style)
    if table:
        t = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def make_xlsx(sheets: dict[str, list[list]]) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class FakeClient:
    """Mimics the parts of anthropic.Anthropic that AnswerGenerator uses.

    `responder(system, user_prompt) -> dict | Exception` decides each reply.
    """

    def __init__(self, responder):
        self.responder = responder
        self.calls: list[dict] = []
        self.messages = SimpleNamespace(create=self._create)
        self.beta = SimpleNamespace(messages=SimpleNamespace(create=self._create))

    def _create(self, **kwargs):
        self.calls.append(kwargs)
        result = self.responder(kwargs["system"], kwargs["messages"][0]["content"])
        if isinstance(result, Exception):
            raise result
        stop = "end_turn"
        if isinstance(result, tuple):
            stop, result = result
        text = result if isinstance(result, str) else json.dumps(result)
        return SimpleNamespace(stop_reason=stop, content=[SimpleNamespace(type="text", text=text)])
