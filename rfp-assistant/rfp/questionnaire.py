"""Find the questions in an uploaded questionnaire."""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

QUESTION_HEADERS = ("question", "questions", "requirement", "requirements", "query", "criteria",
                    "criterion", "request", "description", "item")
ANSWER_HEADERS = ("answer", "response", "vendor response", "supplier response", "reply",
                  "comment", "comments", "vendor comments")
SCAN_ROWS = 20


@dataclass
class Question:
    id: str
    text: str
    kind: str  # "excel", "paragraph" or "table"
    location: dict = field(default_factory=dict)


def looks_like_question(text: str) -> bool:
    text = text.strip()
    if "?" in text:
        return True
    # Imperative requirement lines: "Describe ...", "Provide ...", "Please list ..."
    return len(text) >= 15 and len(text.split()) >= 3


# ===================================================================== Excel


@dataclass
class ExcelLayout:
    sheet: str
    header_row: int  # 0 when there is no header row
    question_col: int  # 1-based
    answer_col: int  # 1-based


def _load(data: bytes, read_only: bool = True):
    return load_workbook(io.BytesIO(data), read_only=read_only, data_only=True)


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def excel_preview(data: bytes, rows: int = 30) -> dict[str, list[list[str]]]:
    wb = _load(data)
    preview = {}
    for ws in wb.worksheets:
        preview[ws.title] = [
            [_text(v) for v in row] for row in ws.iter_rows(max_row=rows, values_only=True)
        ]
    wb.close()
    return preview


def column_label(index: int, header: str = "") -> str:
    letter = get_column_letter(index)
    return f"{letter} – {header}" if header else letter


def _header_match(text: str, keywords: tuple[str, ...]) -> int:
    """2 = exact keyword, 1 = keyword is a word in the text, 0 = no match."""
    lowered = text.lower().strip(" :*")
    if lowered in keywords:
        return 2
    words = set(re.findall(r"[a-z]+", lowered))
    return 1 if words & set(keywords) else 0


def detect_excel_layout(data: bytes) -> ExcelLayout:
    """Guess the sheet, header row, question column and answer column."""
    best: tuple[float, ExcelLayout] | None = None
    wb = _load(data)
    for ws in wb.worksheets:
        rows = [[_text(v) for v in row] for row in ws.iter_rows(values_only=True)]
        if not any(any(r) for r in rows):
            continue
        width = max(len(r) for r in rows)
        rows = [r + [""] * (width - len(r)) for r in rows]

        layout, score = _layout_from_header(ws.title, rows)
        if layout is None:
            layout, score = _layout_from_content(ws.title, rows)
        if layout and (best is None or score > best[0]):
            best = (score, layout)
    wb.close()
    if best is None:
        raise ValueError("The workbook has no text to read questions from.")
    return best[1]


def _layout_from_header(title: str, rows: list[list[str]]):
    """Pick the best header row in the first rows of the sheet. A real header
    row has several filled cells (a title row like 'Vendor Requirements' does not)."""
    best = None
    for r, row in enumerate(rows[:SCAN_ROWS], start=1):
        filled = sum(1 for v in row if v)
        for c, value in enumerate(row, start=1):
            strength = _header_match(value, QUESTION_HEADERS) if value else 0
            if not strength:
                continue
            body = [row_[c - 1] for row_ in rows[r:]]
            count = sum(1 for t in body if t and looks_like_question(t))
            if count == 0:
                continue
            score = (filled >= 2, strength, count, -r)
            if best is None or score > best[0]:
                best = (score, r, c)
    if best is None:
        return None, 0
    (_, strength, count, _), r, q_col = best
    answer_col = _find_answer_col(rows[r - 1], q_col)
    if answer_col is None:
        answer_col = _first_empty_col(rows, r, q_col)
    return ExcelLayout(title, r, q_col, answer_col), 100 * strength + count


def _find_answer_col(header: list[str], q_col: int) -> int | None:
    candidates = [(c, _header_match(v, ANSWER_HEADERS)) for c, v in enumerate(header, start=1)
                  if v and c != q_col]
    candidates = [(c, s) for c, s in candidates if s]
    if not candidates:
        return None
    # Prefer the strongest match, then the closest column to the right of the question.
    candidates.sort(key=lambda cs: (-cs[1], cs[0] < q_col, abs(cs[0] - q_col)))
    return candidates[0][0]


def _first_empty_col(rows: list[list[str]], header_row: int, q_col: int) -> int:
    width = len(rows[0]) if rows else q_col
    for c in range(q_col + 1, width + 1):
        header = rows[header_row - 1][c - 1] if header_row else ""
        column = [row[c - 1] for row in rows[header_row:]]
        if not header and not any(column):
            return c
    return width + 1


def _layout_from_content(title: str, rows: list[list[str]]):
    """No header found: pick the column that looks most like questions."""
    width = len(rows[0])
    best_col, best_score = 0, 0.0
    for c in range(1, width + 1):
        texts = [row[c - 1] for row in rows if row[c - 1]]
        score = sum(2 if "?" in t else (1 if looks_like_question(t) else 0) for t in texts)
        if score > best_score:
            best_col, best_score = c, score
    if not best_col:
        return None, 0
    return ExcelLayout(title, 0, best_col, _first_empty_col(rows, 0, best_col)), best_score


def extract_excel_questions(data: bytes, layout: ExcelLayout) -> list[Question]:
    wb = _load(data)
    if layout.sheet not in wb.sheetnames:
        raise ValueError(f"Sheet {layout.sheet!r} not found")
    ws = wb[layout.sheet]
    questions = []
    for r, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if r <= layout.header_row:
            continue
        text = _text(row[layout.question_col - 1]) if len(row) >= layout.question_col else ""
        if text and looks_like_question(text):
            cell = f"{get_column_letter(layout.question_col)}{r}"
            questions.append(
                Question(id=f"{layout.sheet}!{cell}", text=text, kind="excel",
                         location={"sheet": layout.sheet, "row": r})
            )
    wb.close()
    return questions
