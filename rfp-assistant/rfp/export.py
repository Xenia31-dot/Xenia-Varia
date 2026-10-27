"""Excel export of the review table, and write-back into the original template."""
from __future__ import annotations

import io

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .prompts import NO_ANSWER
from .questionnaire import ExcelLayout

REVIEW_COLUMNS = ["Question", "Draft answer", "Sources", "Confidence", "Status"]
STATUSES = ["Draft", "Approved", "Needs SME"]

_STATUS_FILL = {
    "Approved": PatternFill("solid", fgColor="E2F0D9"),
    "Needs SME": PatternFill("solid", fgColor="FCE4D6"),
}
_WIDTHS = {"Question": 60, "Draft answer": 80, "Sources": 40, "Confidence": 12, "Status": 12}


def review_table_to_excel(table: pd.DataFrame) -> bytes:
    columns = [c for c in REVIEW_COLUMNS if c in table.columns]
    columns += [c for c in table.columns if c not in REVIEW_COLUMNS and not c.startswith("_")]
    out = io.BytesIO()
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        table[columns].to_excel(writer, sheet_name="Answers", index=False)
        ws = writer.sheets["Answers"]
        for c, name in enumerate(columns, start=1):
            ws.column_dimensions[get_column_letter(c)].width = _WIDTHS.get(name, 20)
            ws.cell(1, c).font = Font(bold=True)
        status_col = columns.index("Status") + 1 if "Status" in columns else None
        for r in range(2, ws.max_row + 1):
            for c in range(1, len(columns) + 1):
                ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")
            if status_col:
                fill = _STATUS_FILL.get(ws.cell(r, status_col).value)
                if fill:
                    for c in range(1, len(columns) + 1):
                        ws.cell(r, c).fill = fill
        ws.freeze_panes = "A2"
    return out.getvalue()


def write_back_excel(
    template: bytes,
    layout: ExcelLayout,
    answers: dict[int, str],
    skip_no_answer: bool = False,
    header: str = "Answer",
) -> bytes:
    """Put each answer in `layout.answer_col` on the question's row.

    `answers` maps the 1-based row number to the answer text. The rest of the
    workbook (other sheets, formatting, formulas) is left as it was.
    """
    keep_vba = template[:2] == b"PK" and b"vbaProject" in template
    wb = load_workbook(io.BytesIO(template), keep_vba=keep_vba)
    ws = wb[layout.sheet]
    if layout.header_row and not ws.cell(layout.header_row, layout.answer_col).value:
        ws.cell(layout.header_row, layout.answer_col).value = header
    letter = get_column_letter(layout.answer_col)
    if (ws.column_dimensions[letter].width or 0) < 40:
        ws.column_dimensions[letter].width = 60
    for row, answer in answers.items():
        if skip_no_answer and (not answer or answer.strip() == NO_ANSWER):
            continue
        cell = ws.cell(row, layout.answer_col)
        cell.value = answer
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
