"""Generate sample files for a fictional company, "Nimbus Analytics".

    python -m eval.make_samples

Writes a small knowledge base (PDF, Word, Excel) to eval/sample_docs/ and two
sample questionnaires (Excel and Word) to eval/sample_questionnaires/.
"""
from __future__ import annotations

from pathlib import Path

import docx
from openpyxl import Workbook
from openpyxl.styles import Font

HERE = Path(__file__).resolve().parent
DOCS = HERE / "sample_docs"
QUESTIONNAIRES = HERE / "sample_questionnaires"

SECURITY_PDF_PAGES = [
    [
        "Nimbus Analytics - Security Whitepaper",
        "",
        "Encryption",
        "All customer data is encrypted at rest using AES-256.",
        "Data in transit is protected with TLS 1.2 or higher.",
        "Encryption keys are managed in AWS KMS and rotated every 90 days.",
        "",
        "Identity and access",
        "Nimbus supports single sign-on (SSO) via SAML 2.0 and OpenID Connect.",
        "Multi-factor authentication (MFA) can be enforced for all users.",
        "Role-based access control offers Admin, Editor and Viewer roles.",
    ],
    [
        "Certifications and testing",
        "Nimbus Analytics holds SOC 2 Type II and ISO 27001 certifications.",
        "An independent third party performs a penetration test every year.",
        "Nimbus is not FedRAMP authorized.",
        "",
        "Business continuity",
        "Backups are taken daily and retained for 35 days.",
        "Recovery point objective (RPO) is 24 hours.",
        "Recovery time objective (RTO) is 4 hours.",
        "",
        "Data retention",
        "Customer data is deleted within 30 days after the contract ends.",
    ],
]


def write_pdf(path: Path, pages: list[list[str]]) -> None:
    """Minimal PDF writer (Helvetica, one text line per entry). No extra deps."""
    objects: list[bytes] = []
    page_ids = []
    font_id = 3 + 2 * len(pages)
    for i, lines in enumerate(pages):
        page_id, content_id = 3 + 2 * i, 4 + 2 * i
        page_ids.append(page_id)
        ops = ["BT /F1 11 Tf 14 TL 60 760 Td"]
        for line in lines:
            safe = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            ops.append(f"({safe}) Tj T*")
        ops.append("ET")
        stream = "\n".join(ops).encode("latin-1")
        objects.append(
            f"{page_id} 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Contents {content_id} 0 R /Resources << /Font << /F1 {font_id} 0 R >> >> >> endobj\n".encode()
        )
        objects.append(f"{content_id} 0 obj << /Length {len(stream)} >> stream\n".encode()
                       + stream + b"\nendstream endobj\n")
    kids = " ".join(f"{p} 0 R" for p in page_ids)
    all_objects = [
        b"1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj\n",
        f"2 0 obj << /Type /Pages /Kids [{kids}] /Count {len(pages)} >> endobj\n".encode(),
        *objects,
        f"{font_id} 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj\n".encode(),
    ]
    body = b"%PDF-1.4\n"
    offsets = []
    for obj in all_objects:
        offsets.append(len(body))
        body += obj
    xref = len(body)
    body += f"xref\n0 {len(offsets) + 1}\n0000000000 65535 f \n".encode()
    body += b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    body += f"trailer << /Size {len(offsets) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(body)


def write_product_overview(path: Path) -> None:
    d = docx.Document()
    d.add_heading("Nimbus Analytics - Product Overview", 0)
    d.add_heading("What Nimbus does", 1)
    d.add_paragraph(
        "Nimbus Analytics is a cloud analytics platform that lets business teams build "
        "dashboards and reports without writing code."
    )
    d.add_heading("Hosting and data residency", 1)
    d.add_paragraph(
        "Nimbus is hosted on Amazon Web Services (AWS). Customers choose their data region at "
        "sign-up: EU (Frankfurt, eu-central-1) or US (N. Virginia, us-east-1). Data does not "
        "leave the chosen region. Nimbus is offered as SaaS only; on-premise deployment is not offered."
    )
    d.add_heading("Privacy", 1)
    d.add_paragraph(
        "Nimbus is GDPR compliant and provides a Data Processing Agreement (DPA) to all customers."
    )
    d.add_heading("Integrations", 1)
    d.add_paragraph("Native connectors are available for:")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Connector", "Type"
    for name, kind in [("Snowflake", "Data warehouse"), ("Google BigQuery", "Data warehouse"),
                       ("Salesforce", "CRM"), ("Slack", "Notifications")]:
        row = table.add_row().cells
        row[0].text, row[1].text = name, kind
    d.add_heading("API", 1)
    d.add_paragraph(
        "A REST API covers all platform features. The API rate limit is 1,000 requests per "
        "minute per workspace. Official SDKs are available for Python and JavaScript."
    )
    d.save(path)


def write_service_sheet(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "SLA and Support"
    ws.append(["Item", "Standard plan", "Enterprise plan"])
    for row in [
        ["Uptime SLA", "99.9%", "99.95%"],
        ["Support hours", "Business hours (Mon-Fri 9:00-18:00 CET)", "24/7"],
        ["P1 incident first response", "4 hours", "1 hour"],
        ["Dedicated customer success manager", "No", "Yes"],
    ]:
        ws.append(row)
    prices = wb.create_sheet("Pricing")
    prices.append(["Plan", "Price per user per month (USD)", "Minimum users"])
    prices.append(["Standard", 25, 5])
    prices.append(["Enterprise", 60, 50])
    wb.save(path)


QUESTIONS = [
    ("SEC-1", "Security", "Is customer data encrypted at rest and in transit?"),
    ("SEC-2", "Security", "Does the platform support single sign-on (SSO)? Which protocols?"),
    ("SEC-3", "Security", "Which security certifications does your company hold?"),
    ("SEC-4", "Security", "Are you FedRAMP authorized?"),
    ("SEC-5", "Security", "How often are penetration tests performed?"),
    ("OPS-1", "Operations", "What uptime SLA do you offer for the Enterprise plan?"),
    ("OPS-2", "Operations", "Describe your backup policy, including RPO and RTO."),
    ("OPS-3", "Operations", "What are your support hours?"),
    ("DAT-1", "Data", "Where is customer data hosted?"),
    ("DAT-2", "Data", "Can the solution be deployed on-premise?"),
    ("INT-1", "Integrations", "Which data warehouses can you connect to?"),
    ("COM-1", "Commercial", "What is your company's annual revenue?"),
]


def write_questionnaire_xlsx(path: Path) -> None:
    wb = Workbook()
    cover = wb.active
    cover.title = "Instructions"
    cover["A1"] = "Request for Proposal - Analytics Platform"
    cover["A2"] = "Please answer every question in the 'Questions' sheet."
    ws = wb.create_sheet("Questions")
    ws.append(["Supplier Questionnaire"])
    ws.append([])
    ws.append(["Ref", "Category", "Question", "Supplier Response"])
    for c in range(1, 5):
        ws.cell(3, c).font = Font(bold=True)
    for row in QUESTIONS:
        ws.append(list(row))
    ws.column_dimensions["C"].width = 70
    ws.column_dimensions["D"].width = 60
    wb.save(path)


def write_questionnaire_docx(path: Path) -> None:
    d = docx.Document()
    d.add_heading("Vendor Due Diligence Questionnaire", 0)
    d.add_paragraph("Please answer the questions below and return this document.")
    d.add_heading("Part A - General", 1)
    d.add_paragraph("1. Where is customer data hosted?")
    d.add_paragraph("2. Describe your backup policy, including RPO and RTO.")
    d.add_paragraph("3. What is your company's annual revenue?")
    d.add_heading("Part B - Security", 1)
    table = d.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Question", "Answer"
    for q in ["Is customer data encrypted at rest and in transit?",
              "Does the platform support single sign-on (SSO)?",
              "Are you FedRAMP authorized?"]:
        row = table.add_row().cells
        row[0].text = q
    d.save(path)


def main() -> None:
    DOCS.mkdir(exist_ok=True)
    QUESTIONNAIRES.mkdir(exist_ok=True)
    write_pdf(DOCS / "Nimbus Security Whitepaper.pdf", SECURITY_PDF_PAGES)
    write_product_overview(DOCS / "Nimbus Product Overview.docx")
    write_service_sheet(DOCS / "Nimbus Service Levels and Pricing.xlsx")
    write_questionnaire_xlsx(QUESTIONNAIRES / "Sample RFP.xlsx")
    write_questionnaire_docx(QUESTIONNAIRES / "Sample Due Diligence.docx")
    print(f"Wrote sample documents to {DOCS} and questionnaires to {QUESTIONNAIRES}")


if __name__ == "__main__":
    main()
